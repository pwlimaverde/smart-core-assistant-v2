"""Leitura do texto de documentos de treinamento (B9 / N10 E5).

Camada técnica ("burra"): recebe bytes, devolve texto, e falha lançando exceção
— a tradução para erro de domínio é do repositório.

Formatos aceitos: pdf, docx, xlsx, txt e csv. Os binários antigos do Office
(`.doc`, `.xls`) ficam de fora de propósito: lê-los exige conversores externos
(LibreOffice, antiword) que não cabem na imagem do serviço, e a saída de
"salve como .docx/.xlsx" é trivial para quem tem o arquivo.

Nada aqui loga o texto: documento de treinamento pode ter dado sensível do
negócio.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable

MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_POR_MIMETYPE = {
    MIME_PDF: "pdf",
    MIME_DOCX: "docx",
    MIME_XLSX: "xlsx",
    "text/plain": "txt",
    "text/csv": "csv",
}
_POR_EXTENSAO = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".txt": "txt",
    ".csv": "csv",
}

MAX_PAGINAS_PADRAO = 300
MAX_CARACTERES_PADRAO = 500_000


class FormatoNaoSuportadoException(Exception):
    """O arquivo não é de um formato que sabemos ler."""


class DocumentoIlegivelException(Exception):
    """O arquivo é do formato certo, mas não dá para ler (corrompido, senha)."""


def formato_de(mimetype: str, nome_arquivo: str) -> str | None:
    """Formato pelo mimetype; sem mimetype conhecido, pela extensão do nome."""
    base = (mimetype or "").split(";")[0].strip().lower()
    if base in _POR_MIMETYPE:
        return _POR_MIMETYPE[base]
    nome = (nome_arquivo or "").lower()
    for extensao, formato in _POR_EXTENSAO.items():
        if nome.endswith(extensao):
            return formato
    return None


def extrair(
    conteudo: bytes,
    mimetype: str,
    nome_arquivo: str,
    *,
    max_paginas: int = MAX_PAGINAS_PADRAO,
    max_caracteres: int = MAX_CARACTERES_PADRAO,
) -> tuple[str, str]:
    """Devolve `(texto, formato)`.

    O texto é truncado em `max_caracteres`: um documento gigante vira milhares
    de trechos vetorizados, e o custo disso é do tenant.

    Raises:
        FormatoNaoSuportadoException: formato fora da lista.
        DocumentoIlegivelException: arquivo corrompido ou protegido por senha.
    """
    formato = formato_de(mimetype, nome_arquivo)
    if formato is None:
        raise FormatoNaoSuportadoException(
            "formato não suportado; envie pdf, docx, xlsx, txt ou csv "
            "(arquivos .doc e .xls: salve como .docx ou .xlsx)"
        )
    leitor: Callable[[bytes, int], str] = _LEITORES[formato]
    try:
        texto = leitor(conteudo, max_paginas)
    except DocumentoIlegivelException:
        raise
    except Exception as exc:  # noqa: BLE001 — cada biblioteca falha do seu jeito
        raise DocumentoIlegivelException(
            f"não foi possível ler o arquivo {formato}: {type(exc).__name__}"
        ) from exc
    return texto.strip()[:max_caracteres], formato


def _ler_pdf(conteudo: bytes, max_paginas: int) -> str:
    from pypdf import PdfReader

    leitor = PdfReader(io.BytesIO(conteudo))
    if leitor.is_encrypted:
        raise DocumentoIlegivelException(
            "o PDF está protegido por senha; envie uma cópia sem senha"
        )
    paginas = [pagina.extract_text() or "" for pagina in leitor.pages[:max_paginas]]
    return "\n\n".join(p.strip() for p in paginas if p.strip())


def _ler_docx(conteudo: bytes, _max_paginas: int) -> str:
    import docx

    documento = docx.Document(io.BytesIO(conteudo))
    partes = [p.text.strip() for p in documento.paragraphs if p.text.strip()]
    # Tabelas guardam boa parte do conteúdo útil (preços, horários): ler só os
    # parágrafos as perderia em silêncio.
    for tabela in documento.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))
    return "\n".join(partes)


def _ler_xlsx(conteudo: bytes, _max_paginas: int) -> str:
    from openpyxl import load_workbook

    planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    partes: list[str] = []
    try:
        for aba in planilha.worksheets:
            linhas = [
                " | ".join(str(v).strip() for v in linha if v is not None)
                for linha in aba.iter_rows(values_only=True)
            ]
            linhas = [linha for linha in linhas if linha.strip()]
            if linhas:
                partes.append(f"## {aba.title}")
                partes.extend(linhas)
    finally:
        planilha.close()
    return "\n".join(partes)


def _decodificar(conteudo: bytes) -> str:
    # UTF-8 primeiro (com BOM, que o Excel grava); Latin-1 nunca falha, e é o
    # que sobra de planilhas antigas exportadas no Windows.
    try:
        return conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        return conteudo.decode("latin-1")


def _ler_txt(conteudo: bytes, _max_paginas: int) -> str:
    return _decodificar(conteudo)


def _ler_csv(conteudo: bytes, _max_paginas: int) -> str:
    texto = _decodificar(conteudo)
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        delimitador = dialeto.delimiter
    except csv.Error:
        delimitador = ","
    linhas = [
        " | ".join(celula.strip() for celula in linha if celula.strip())
        for linha in csv.reader(io.StringIO(texto), delimiter=delimitador)
    ]
    return "\n".join(linha for linha in linhas if linha)


_LEITORES: dict[str, Callable[[bytes, int], str]] = {
    "pdf": _ler_pdf,
    "docx": _ler_docx,
    "xlsx": _ler_xlsx,
    "txt": _ler_txt,
    "csv": _ler_csv,
}
