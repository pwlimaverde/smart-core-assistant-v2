"""Testes da feature ExtrairTextoDocumento (B9 / N10 E5).

Os documentos são gerados em memória pelas próprias bibliotecas de leitura
(python-docx, openpyxl, pypdf) — nada de arquivo binário versionado. O download
é trocado por um fake: nenhum teste toca rede.
"""

from __future__ import annotations

import io

import pytest
from py_return_success_or_error import ErrorGeneric, Failure, Success

from ia_engine.domain.errors import MediaDownloadError
from ia_engine.features.extrair_texto import (
    DocumentoIlegivelError,
    ExtrairTextoDataSource,
    ExtrairTextoParameters,
    ExtrairTextoRepository,
    ExtrairTextoUsecase,
    FormatoNaoSuportadoError,
    TextoExtraido,
    TextoVazioError,
)
from ia_engine.features.extrair_texto.datasources import (
    extrair_texto_datasource as modulo_datasource,
)
from ia_engine.features.extrair_texto.services.leitores import (
    MIME_DOCX,
    MIME_PDF,
    MIME_XLSX,
    DocumentoIlegivelException,
    FormatoNaoSuportadoException,
    extrair,
    formato_de,
)
from ia_engine.shared.media import MediaDownloadException


# ------------------------------------------------------------ geradores
def _pdf_com_texto(texto: str) -> bytes:
    """PDF mínimo, de uma página, com o texto em Helvetica."""
    fluxo = f"BT /F1 12 Tf 72 712 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(fluxo)).encode()
        + b" >>\nstream\n"
        + fluxo
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    saida = io.BytesIO()
    saida.write(b"%PDF-1.4\n")
    posicoes = []
    for numero, corpo in enumerate(objetos, start=1):
        posicoes.append(saida.tell())
        saida.write(f"{numero} 0 obj\n".encode() + corpo + b"\nendobj\n")
    xref = saida.tell()
    saida.write(f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode())
    for posicao in posicoes:
        saida.write(f"{posicao:010d} 00000 n \n".encode())
    saida.write(
        f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    return saida.getvalue()


def _pdf_com_senha() -> bytes:
    from pypdf import PdfWriter

    escritor = PdfWriter()
    escritor.add_blank_page(width=200, height=200)
    escritor.encrypt("segredo")
    saida = io.BytesIO()
    escritor.write(saida)
    return saida.getvalue()


def _docx() -> bytes:
    import docx

    documento = docx.Document()
    documento.add_paragraph("Entregamos de segunda a sábado.")
    documento.add_paragraph("   ")
    tabela = documento.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text = "Produto"
    tabela.cell(0, 1).text = "Preço"
    tabela.cell(1, 0).text = "Bolo"
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    planilha = Workbook()
    aba = planilha.active
    assert aba is not None
    aba.title = "Precos"
    aba.append(["Produto", "Preço"])
    aba.append(["Bolo", 42])
    aba.append([None, None])
    planilha.create_sheet("Vazia")
    saida = io.BytesIO()
    planilha.save(saida)
    return saida.getvalue()


# ------------------------------------------------------------ leitores
def test_formato_pelo_mimetype_ou_pela_extensao() -> None:
    assert formato_de("application/pdf; charset=binary", "x") == "pdf"
    assert formato_de("", "Tabela.XLSX") == "xlsx"
    assert formato_de("application/octet-stream", "notas.txt") == "txt"
    assert formato_de("application/msword", "antigo.doc") is None


def test_le_pdf() -> None:
    texto, formato = extrair(_pdf_com_texto("Ola mundo"), MIME_PDF, "a.pdf")
    assert formato == "pdf"
    assert "Ola mundo" in texto


def test_pdf_com_senha_e_ilegivel() -> None:
    with pytest.raises(DocumentoIlegivelException, match="senha"):
        extrair(_pdf_com_senha(), MIME_PDF, "a.pdf")


def test_pdf_corrompido_e_ilegivel() -> None:
    with pytest.raises(DocumentoIlegivelException, match="pdf"):
        extrair(b"%PDF-lixo", MIME_PDF, "a.pdf")


def test_le_docx_com_tabela() -> None:
    texto, formato = extrair(_docx(), MIME_DOCX, "a.docx")
    assert formato == "docx"
    assert "Entregamos de segunda a sábado." in texto
    assert "Produto | Preço" in texto
    assert "Bolo" in texto


def test_le_xlsx_por_aba_e_ignora_aba_vazia() -> None:
    texto, formato = extrair(_xlsx(), MIME_XLSX, "a.xlsx")
    assert formato == "xlsx"
    assert "## Precos" in texto
    assert "Bolo | 42" in texto
    assert "Vazia" not in texto


def test_le_txt_em_utf8_e_latin1() -> None:
    assert extrair("olá".encode(), "text/plain", "a.txt")[0] == "olá"
    assert extrair("olá".encode("latin-1"), "text/plain", "a.txt")[0] == "olá"


def test_le_csv_com_ponto_e_virgula() -> None:
    texto, formato = extrair(
        "produto;preço\nbolo;42\n;\n".encode(), "text/csv", "a.csv"
    )
    assert formato == "csv"
    assert texto == "produto | preço\nbolo | 42"


def test_csv_sem_delimitador_reconhecivel_usa_virgula() -> None:
    assert extrair(b"linha unica", "text/csv", "a.csv")[0] == "linha unica"


def test_formato_antigo_do_office_e_recusado() -> None:
    with pytest.raises(FormatoNaoSuportadoException, match="docx"):
        extrair(b"\xd0\xcf\x11\xe0", "application/msword", "a.doc")


def test_texto_e_truncado() -> None:
    texto, _ = extrair(b"x" * 50, "text/plain", "a.txt", max_caracteres=10)
    assert texto == "x" * 10


# ------------------------------------------------------------ cadeia RSOE
def _usecase() -> ExtrairTextoUsecase:
    return ExtrairTextoUsecase(ExtrairTextoRepository(ExtrairTextoDataSource()))


def _parametros(mimetype: str = "text/plain") -> ExtrairTextoParameters:
    return ExtrairTextoParameters(
        url="https://r2/assinada", mimetype=mimetype, nome_arquivo="a.txt"
    )


def _baixa(monkeypatch: pytest.MonkeyPatch, resultado: bytes | Exception) -> None:
    async def _fake(url: str, **_kwargs: object) -> bytes:
        if isinstance(resultado, Exception):
            raise resultado
        return resultado

    monkeypatch.setattr(modulo_datasource, "download_media", _fake)


async def test_extrai_o_texto(monkeypatch: pytest.MonkeyPatch) -> None:
    _baixa(monkeypatch, b"  Horario: 8h as 18h  ")
    resultado = await _usecase()(_parametros())
    match resultado:
        case Success(valor):
            assert valor == TextoExtraido(texto="Horario: 8h as 18h", formato="txt")
        case _:
            pytest.fail(f"esperava sucesso, veio {resultado}")


async def test_arquivo_sem_texto_e_falha(monkeypatch: pytest.MonkeyPatch) -> None:
    _baixa(monkeypatch, b"   ")
    resultado = await _usecase()(_parametros())
    assert isinstance(resultado, Failure)
    assert isinstance(resultado.error, TextoVazioError)


@pytest.mark.parametrize(
    ("falha", "mimetype", "esperado"),
    [
        (MediaDownloadException("404"), "text/plain", MediaDownloadError),
        (None, "application/msword", FormatoNaoSuportadoError),
        (None, MIME_PDF, DocumentoIlegivelError),
        (RuntimeError("inesperado"), "text/plain", ErrorGeneric),
    ],
)
async def test_falhas_viram_erro_de_dominio(
    monkeypatch: pytest.MonkeyPatch,
    falha: Exception | None,
    mimetype: str,
    esperado: type,
) -> None:
    _baixa(monkeypatch, falha if falha is not None else b"%PDF-lixo")
    resultado = await _usecase()(_parametros(mimetype))
    assert isinstance(resultado, Failure)
    assert isinstance(resultado.error, esperado)


def test_usecase_converte_excecao_fora_da_fronteira() -> None:
    erro = _usecase().on_unexpected(ValueError("x"))
    assert isinstance(erro, ErrorGeneric)
    assert "ValueError" in erro.message
